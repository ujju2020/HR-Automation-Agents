import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def create_deck():
    template_path = r"C:\Users\Administration\Downloads\Idea Submission Deck _ Builders Pitch Fest 2026.pptx"
    output_pptx = r"c:\Users\Administration\Desktop\VibeCoding\BITSom Vertex Fest\HR-Automation-Agents\BPF2026_IdeaSubmission_PitchDeck_PulseHR.pptx"
    output_pptx_dl = r"C:\Users\Administration\Downloads\BPF2026_IdeaSubmission_PitchDeck_PulseHR.pptx"

    prs = Presentation(template_path)

    # Color Palette (HR Tech & Modern Swarm)
    DARK_NAVY = RGBColor(15, 23, 42)      # #0F172A - Headers
    ROYAL_INDIGO = RGBColor(99, 102, 241) # #6366F1 - Question headers / Accents
    SLATE_BODY = RGBColor(51, 65, 85)     # #334155 - Body text
    MUTED_TEXT = RGBColor(100, 116, 139)  # #64748B - Sub-meta
    CYAN_ACCENT = RGBColor(14, 165, 233)  # #0EA5E9 - Metrics / highlights

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide (Cover)
    # -------------------------------------------------------------
    slide1 = prs.slides[0]
    tb1 = slide1.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(8.4), Inches(3.2))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "PulseHR OS"
    p0.font.name = "Arial"
    p0.font.size = Pt(36)
    p0.font.bold = True
    p0.font.color.rgb = ROYAL_INDIGO
    p0.space_after = Pt(6)

    p1 = tf1.add_paragraph()
    p1.text = "Autonomous Multi-Agent Human Capital Management & Talent Lifecycle Platform"
    p1.font.name = "Arial"
    p1.font.size = Pt(17)
    p1.font.bold = True
    p1.font.color.rgb = DARK_NAVY
    p1.space_after = Pt(12)

    p2 = tf1.add_paragraph()
    p2.text = "5-Agent Autonomous Swarm • Bias-Mitigated Screening • Real-Time Burnout & Retention Telemetry"
    p2.font.name = "Arial"
    p2.font.size = Pt(11.5)
    p2.font.color.rgb = SLATE_BODY
    p2.space_after = Pt(20)

    p3 = tf1.add_paragraph()
    p3.text = "BITSoM Vertex Pitch Fest 2026 (BPF 2026) | Idea Submission Deck"
    p3.font.name = "Arial"
    p3.font.size = Pt(12)
    p3.font.bold = True
    p3.font.color.rgb = CYAN_ACCENT
    p3.space_after = Pt(4)

    p4 = tf1.add_paragraph()
    p4.text = "Founder: Ujjwal Kumar Bhowmick  |  Email: ujjwalkumarbhowmick30@gmail.com"
    p4.font.name = "Arial"
    p4.font.size = Pt(11)
    p4.font.color.rgb = MUTED_TEXT

    # -------------------------------------------------------------
    # Helper to populate Q&A content into Shape on slides 3-13
    # -------------------------------------------------------------
    def populate_qa_slide(slide_idx, qa_list):
        slide = prs.slides[slide_idx - 1]
        target_shape = None
        for s in slide.shapes:
            if s.has_text_frame and s.top > 1000000:
                target_shape = s
                break
        
        if not target_shape:
            target_shape = slide.shapes.add_textbox(Inches(0.4), Inches(1.3), Inches(9.2), Inches(3.6))

        tf = target_shape.text_frame
        tf.word_wrap = True
        tf.clear()

        first = True
        for item in qa_list:
            if "q" in item:
                pq = tf.paragraphs[0] if first else tf.add_paragraph()
                first = False
                pq.text = item["q"]
                pq.font.name = "Arial"
                pq.font.size = Pt(11)
                pq.font.bold = True
                pq.font.color.rgb = ROYAL_INDIGO
                pq.space_before = Pt(6) if not first else Pt(0)
                pq.space_after = Pt(3)

            for b in item.get("bullets", []):
                pb = tf.add_paragraph()
                pb.text = "•  " + b
                pb.font.name = "Arial"
                pb.font.size = Pt(9.5)
                pb.font.color.rgb = SLATE_BODY
                pb.space_after = Pt(2.5)

    # -------------------------------------------------------------
    # SLIDE 3: Startup Snapshot
    # -------------------------------------------------------------
    populate_qa_slide(3, [
        {
            "q": "What problem are you trying to solve?",
            "bullets": [
                "People Operations Fragmentation: Enterprise HR is trapped in disconnected silos (ATS, HRIS, review tools). Teams spend 60% of their bandwidth on manual resume triage, chaotic onboarding checklists, and repetitive policy tickets.",
                "Preventable Silent Attrition: Companies lose top performers to silent burnout because organizational sentiment is assessed reactively via annual surveys rather than through continuous early-warning telemetry."
            ]
        },
        {
            "q": "What inspired you to work on this problem?",
            "bullets": [
                "The Multi-Agent Swarm Breakthrough: Instead of another static HR portal, employee lifecycle stages require dedicated, specialized autonomous AI agents collaborating as a synchronized team to empower HR leaders and employees.",
                "The Solution: PulseHR OS coordinates 5 specialized agents across screening, onboarding, burnout telemetry, 360 performance reviews, and policy assistance."
            ]
        }
    ])

    # -------------------------------------------------------------
    # SLIDE 4: Problem Understanding
    # -------------------------------------------------------------
    populate_qa_slide(4, [
        {
            "q": "What problem are you solving, and who experiences it most acutely?",
            "bullets": [
                "Acutely experienced by Chief People Officers (CPOs), Heads of Talent Acquisition, HRBPs, and People Managers battling high turnover and administrative overload.",
                "Cost of Turnover: Replacing an enterprise employee costs 1.5x to 2x their annual salary. Average time-to-hire exceeds 44 days with significant recruiter friction.",
                "Cognitive Appraisal Bias: 360-degree reviews suffer from recency bias, anecdotal evaluations, and lack of objective multi-source competency telemetry."
            ]
        },
        {
            "q": "What research, observations, or evidence validate this opportunity?",
            "bullets": [
                "Gallup's State of the Global Workplace reveals employee disengagement and burnout drain $8.9 Trillion globally in lost productivity.",
                "Over 76% of HR leaders report administrative overhead prevents them from executing strategic talent retention and leadership development.",
                "Over 30% of new hires leave within 90 days due to unstructured onboarding experiences."
            ]
        }
    ])

    # -------------------------------------------------------------
    # SLIDE 5: Customer & Opportunity
    # -------------------------------------------------------------
    populate_qa_slide(5, [
        {
            "q": "Who is your ideal customer, and who would make the buying decision?",
            "bullets": [
                "Ideal Customers: Mid-to-Large Enterprises (500–10,000 employees) in Technology, BFSI, Healthcare, and Professional Services.",
                "Economic Buyers: Chief Human Resources Officers (CHRO), Chief People Officers (CPO), VP of Talent, and Chief Operating Officers (COO).",
                "End Users: HRBPs, Talent Recruiters, Hiring Managers, Team Leads, and Employees."
            ]
        },
        {
            "q": "How large is the opportunity if your assumptions prove correct?",
            "bullets": [
                "Global TAM: $38.4 Billion HR Tech and Talent Management Software market (growing at 10.4% CAGR).",
                "India & APAC SAM: $4.2 Billion regional addressable spend on recruitment automation, engagement telemetry, and appraisal systems.",
                "Immediate SOM: ₹240 Crores ($29M ARR) targeting high-growth technology companies and enterprise employers across India."
            ]
        }
    ])

    # -------------------------------------------------------------
    # SLIDE 6: Proposed Solution
    # -------------------------------------------------------------
    populate_qa_slide(6, [
        {
            "q": "What is your proposed solution, and how does it address the identified problem?",
            "bullets": [
                "PulseHR OS is an enterprise multi-agent operating system coordinating 5 specialized autonomous agents across recruitment, onboarding, burnout telemetry, 360 performance reviews, and employee policy assistance."
            ]
        },
        {
            "q": "What are the key capabilities you envision?",
            "bullets": [
                "TalentScout Agent: Multi-vector semantic scoring across skills, culture, and retention with blind algorithmic DEI audits to mitigate demographic bias.",
                "OnboardPilot Agent: Automated 30-60-90 day milestone roadmap with mentor matching and IT provisioning workflows.",
                "PulseSentinel Agent: Real-time team sentiment analysis, meeting load fatigue tracking, and proactive managerial burnout alerts.",
                "Elevate360 Agent: Synthesizes multi-stakeholder feedback into 5-axis competency radars and calibrated promotion indices.",
                "PeopleAdvisor Agent: 24/7 conversational employee desk citing company policy handbooks with exact section references."
            ]
        },
        {
            "q": "What measurable value do you expect the solution to deliver?",
            "bullets": [
                "65% Reduction in recruitment screening and interview scheduling cycle times.",
                "40% Faster new hire ramp to full productivity  |  30% Decrease in preventable employee turnover."
            ]
        }
    ])

    # -------------------------------------------------------------
    # SLIDE 7: Technology & AI Approach
    # -------------------------------------------------------------
    populate_qa_slide(7, [
        {
            "q": "What role will AI play in your proposed solution, and why is it essential?",
            "bullets": [
                "AI acts as an autonomous collaborative workforce: synthesizing candidate qualifications, diagnosing organizational sentiment patterns, generating personalized onboarding paths, and interpreting nuanced policy documents."
            ]
        },
        {
            "q": "What technologies, AI models, or frameworks do you plan to use, and why?",
            "bullets": [
                "Agentic Orchestration Core: Multi-agent state machine powered by Google Gemini 2.5 Flash / Vertex AI with streaming responses and structured JSON contracts.",
                "Modern Stack: React 18, TypeScript, Tailwind CSS, Lucide icons, Recharts visualization, and resilient offline simulation fallback."
            ]
        },
        {
            "q": "What makes you feel your product is a genuine AI product, and not an AI wrapper?",
            "bullets": [
                "PulseHR OS coordinates an interconnected multi-agent state graph where each agent produces structured artifacts (competency radars, DEI audits, 30-60-90 roadmaps) passed downstream across the employee lifecycle."
            ]
        }
    ])

    # -------------------------------------------------------------
    # SLIDE 8: Competitive Landscape
    # -------------------------------------------------------------
    populate_qa_slide(8, [
        {
            "q": "What alternatives currently exist?",
            "bullets": [
                "Legacy HRIS Suites (Workday, SAP SuccessFactors): Monolithic, complex, and lack generative multi-agent collaboration or real-time sentiment telemetry.",
                "Point ATS Platforms (Greenhouse, Lever): Limited strictly to applicant tracking without lifecycle continuity into onboarding and appraisal.",
                "Lagging Survey Tools (Culture Amp, Lattice): Rely on periodic quarterly surveys rather than continuous real-time burnout telemetry and proactive playbooks."
            ]
        },
        {
            "q": "Why do you believe your approach is better suited to solving this problem?",
            "bullets": [
                "PulseHR OS delivers an end-to-end, proactive multi-agent collective covering the complete employee lifecycle from first interview to executive promotion."
            ]
        }
    ])

    # -------------------------------------------------------------
    # SLIDE 9: Business Model & Go-To-Market
    # -------------------------------------------------------------
    populate_qa_slide(9, [
        {
            "q": "Who do you expect will pay for your solution?",
            "bullets": [
                "B2B SaaS Subscription (PEPM - Per Employee Per Month):",
                "Growth Tier: $5 / employee / month (for companies with 200–1,000 employees).",
                "Enterprise Tier: $8 / employee / month (includes full 5-agent swarm, custom handbook fine-tuning, and priority integrations).",
                "Custom Enterprise SLA: Private VPC and on-premise deployments for highly regulated BFSI/Healthcare clients."
            ]
        },
        {
            "q": "How do you plan to acquire your first customers?",
            "bullets": [
                "Design Partner Program: Engage 4 mid-market tech companies (300–800 headcount) as pilot partners to benchmark attrition reduction.",
                "CHRO & HRBP Communities: Thought leadership campaigns demonstrating DEI audit and burnout mitigation playbooks.",
                "BITSoM Vertex Network: Engage corporate partners and alumni enterprise leaders for executive sponsor introductions."
            ]
        }
    ])

    # -------------------------------------------------------------
    # SLIDE 10: Development Roadmap
    # -------------------------------------------------------------
    populate_qa_slide(10, [
        {
            "q": "What are the next major milestones toward building your MVP?",
            "bullets": [
                "Current Status (Phase 0): Fully functional, interactive production prototype with all 5 specialized agents, live candidate evaluation sandbox, sentiment telemetry, and policy assistant (Completed & Verified).",
                "Phase 1 (Months 1–3): Native ATS/HRIS bi-directional sync (Workday, BambooHR, Slack, Microsoft Teams).",
                "Phase 2 (Months 4–6): Predictive workforce planning and compensation benchmarking engine.",
                "Phase 3 (Months 7–12): Enterprise SOC2 Type II, GDPR, and ISO 27001 data residency compliance."
            ]
        },
        {
            "q": "What resources or support will be most critical during this journey?",
            "bullets": [
                "Enterprise Design Partnerships: Access to active HR teams for feedback on candidate rubrics and appraisal calibration.",
                "Executive HR Advisory: Guidance from seasoned Chief People Officers on compliance and change management.",
                "Cloud Infrastructure Support: Google Cloud / Vertex AI enterprise scaling credits."
            ]
        }
    ])

    # -------------------------------------------------------------
    # SLIDE 11: Team
    # -------------------------------------------------------------
    populate_qa_slide(11, [
        {
            "q": "Why is your team uniquely positioned to solve this problem?",
            "bullets": [
                "Strong technical capabilities in multi-agent orchestration, full-stack enterprise web development, and human-centric software design.",
                "Demonstrated execution speed in architecting complex, production-grade applications with clean UI and resilient AI pipelines."
            ]
        },
        {
            "q": "What relevant domain or technical expertise does the team possess?",
            "bullets": [
                "Ujjwal Kumar Bhowmick (Founder & Lead AI Engineer):",
                "Specialized in Multi-Agent AI system design, TypeScript, React 18, Tailwind CSS, and Google Gemini / Vertex AI integrations.",
                "Deep understanding of organizational psychology, enterprise HR workflows, and bias mitigation algorithms."
            ]
        }
    ])

    # -------------------------------------------------------------
    # SLIDE 12: Why BITSoM Vertex?
    # -------------------------------------------------------------
    populate_qa_slide(12, [
        {
            "q": "Why have you applied to the BITSoM Vertex programme?",
            "bullets": [
                "Enterprise Networking: BITSoM's premier business network and executive leadership ecosystem provide direct access to corporate CHROs and enterprise HR decision-makers.",
                "GTM & Business Mentorship: Vertex provides the exact strategic framework required to transition an enterprise multi-agent prototype into high-velocity B2B enterprise contracts."
            ]
        },
        {
            "q": "What is the biggest challenge preventing you from building your MVP?",
            "bullets": [
                "The prototype is fully built and operational! The primary objective now is enterprise customer pilots: securing enterprise sandbox access, refining organizational telemetry against live HRIS data, and scaling our GTM funnel. BITSoM Vertex is the ideal partner."
            ]
        }
    ])

    # -------------------------------------------------------------
    # SLIDE 13: Supporting Material
    # -------------------------------------------------------------
    populate_qa_slide(13, [
        {
            "q": "Key Links & Prototype Validation:",
            "bullets": [
                "Working Prototype: Fully functional local environment (React 18 / TypeScript on :5173 with 5-agent swarm, talent screening sandbox, 30-60-90 ramp, burnout radar, and HR policy assistant).",
                "GitHub Repository: https://github.com/ujju2020/HR-Automation-Agents (Complete source code, agent schemas, and documentation).",
                "Demo Explainer Video: 2-minute comprehensive walkthrough of TalentScout, OnboardPilot, PulseSentinel, Elevate360, and PeopleAdvisor in action.",
                "Contact: Ujjwal Kumar Bhowmick | ujjwalkumarbhowmick30@gmail.com | India"
            ]
        }
    ])

    # -------------------------------------------------------------
    # SLIDE 14: Conclusion / Closing Slide
    # -------------------------------------------------------------
    slide14 = prs.slides[13]
    tb14 = slide14.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(8.4), Inches(3.0))
    tf14 = tb14.text_frame
    tf14.word_wrap = True

    c0 = tf14.paragraphs[0]
    c0.text = "PulseHR OS"
    c0.font.name = "Arial"
    c0.font.size = Pt(36)
    c0.font.bold = True
    c0.font.color.rgb = ROYAL_INDIGO
    c0.space_after = Pt(8)

    c1 = tf14.add_paragraph()
    c1.text = "Empowering Enterprise Teams Through Autonomous Multi-Agent Human Capital Intelligence"
    c1.font.name = "Arial"
    c1.font.size = Pt(17)
    c1.font.bold = True
    c1.font.color.rgb = DARK_NAVY
    c1.space_after = Pt(14)

    c2 = tf14.add_paragraph()
    c2.text = "Thank You! Open for Questions & Evaluation."
    c2.font.name = "Arial"
    c2.font.size = Pt(14)
    c2.font.bold = True
    c2.font.color.rgb = CYAN_ACCENT
    c2.space_after = Pt(16)

    c3 = tf14.add_paragraph()
    c3.text = "Founder: Ujjwal Kumar Bhowmick  |  ujjwalkumarbhowmick30@gmail.com"
    c3.font.name = "Arial"
    c3.font.size = Pt(11)
    c3.font.color.rgb = MUTED_TEXT

    # Save PPTX
    prs.save(output_pptx)
    prs.save(output_pptx_dl)
    print(f"PPTX successfully created at:\n  - {output_pptx}\n  - {output_pptx_dl}")

if __name__ == "__main__":
    create_deck()
